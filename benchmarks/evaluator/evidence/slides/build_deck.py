#!/usr/bin/env python3
"""Build the self-contained drperf cases HTML deck and its PPTX twin."""

from __future__ import annotations

import html
import re
from dataclasses import dataclass, field
from pathlib import Path


ROOT = Path("/home/ubuntu/drperf-cases")
SLIDES = ROOT / "slides"
CASES = ROOT / "CASES.md"
HTML_OUT = SLIDES / "drperf_cases.html"
PPTX_OUT = SLIDES / "drperf_cases.pptx"


@dataclass
class Slide:
    title: str
    body: str
    notes: str
    bullets: list[str] = field(default_factory=list)
    mono: list[str] = field(default_factory=list)


FORMULAS: list[str] = []
MARKERS: list[tuple[str, str]] = []


def esc(s: str) -> str:
    return html.escape(s, quote=False)


def mono(lines: list[str], *, formula: bool = False) -> str:
    if formula:
        FORMULAS.extend(lines)
    return '<pre class="mono">' + "".join(
        f'<div class="ln">{esc(line)}</div>' for line in lines
    ) + "</pre>"


def block(kind: str, label: str, content: str) -> str:
    return (
        f'<div class="block b-{kind}"><div class="blabel">{label}</div>'
        f'<div class="bcontent">{content}</div></div>'
    )


def marker(path: str, line: str) -> tuple[str, str]:
    MARKERS.append((path, line))
    source = ROOT / path if not path.startswith("/") else Path(path)
    hits = [i for i, text in enumerate(source.read_text().splitlines(), 1) if line in text]
    lineno = hits[0] if hits else "?"
    content = f'<p class="pr">{esc(path)}:{lineno}</p>{mono([line])}'
    return content, f"{path}:{lineno}\n{line}"


def case_slide(
    title: str,
    marker_path: str,
    marker_line: str,
    states: str,
    formula: str,
    unexplained: str,
    surprise: str,
    attribution: list[str],
    source: str,
    optimise: str,
    optimise_lines: list[str] | None = None,
    notes: str = "",
) -> Slide:
    ann_html, marker_mono = marker(marker_path, marker_line)
    fit = (
        f'<p class="pr"><span class="what">Declared states:</span> {esc(states)}. '
        f'<span class="what">Unexplained:</span> {esc(unexplained)}.</p>'
        + mono([formula], formula=True)
    )
    sur = f'<p class="pr">{esc(surprise)}</p>' + mono(attribution)
    sur += f'<p class="pr">{esc(source)}</p>'
    opt = f'<p class="pr">{esc(optimise)}</p>'
    if optimise_lines:
        opt += mono(optimise_lines, formula=False)
    body = '<div class="body case">' + block("annotation", "MARKER", ann_html)
    body += block("fit", "FIT", fit)
    body += block("surprise", "SURPRISE", sur)
    body += block("optimise", "OPTIMISE", opt) + "</div>"
    bullets = [
        f"FIT — {states}; {unexplained} unexplained.",
        f"SURPRISE — {surprise} {source}",
        f"OPTIMISE — {optimise}",
    ]
    return Slide(title, body, notes, bullets, [marker_mono, formula] + (optimise_lines or []))


slides: list[Slide] = []


title_body = (
    '<div class="titleslide"><h1>drperf: cost functions and relations, twenty-one cases</h1>'
    '<p class="sub">ditto-kv FTL, stock vLLM 0.28 CPU, and Wan 2.1 in diffusers</p>'
    '<div class="rule"></div>'
    '<p class="meta">Instruction counts per call, exact and reproducible, as a function of the states each region declares,</p>'
    '<p class="meta">with every coefficient attributed to the function that produced it, and exact relations learned between regions.</p></div>'
)
slides.append(Slide(
    "drperf: cost functions and relations, twenty-one cases",
    title_body,
    "Eight FTL cases, nine vLLM cases, and four Wan cases. vLLM A and B each have two measured third beats, so each gets two slides.",
    ["Per-region affine instruction-count models", "Every coefficient attributed to functions", "Exact relations learned between region states"],
    ["cost(state) = coefficient*state + constant"],
))

outputs_body = '''<div class="body">
<div class="block"><div class="blabel">A COST FUNCTION PER REGION, EVERY COEFFICIENT ATTRIBUTED TO A FUNCTION</div>
<pre class="mono"><div class="ln">derive case_admission_limit</div><div class="ln">  cost(superblocks) = 23.3*superblocks + 3,330.7        [all superblocks]   </div><div class="ln cont">blocks: 6 affine, 924 constant, 15 irregular (2.2% of cost)</div><div class="ln">    per-superblocks coefficient by function:</div><div class="ln">                23.3  &lt;_ditto_ftl_core::directory::Directory&gt;::stats  [_ditto_ftl_core.abi3.so]</div></pre>
</div>
<div class="block"><div class="blabel">AND A RELATION BETWEEN REGIONS, LEARNED, EXACT AT EVERY TRIGGER</div>
<pre class="mono"><div class="ln">attention.kv_tokens = count(detokenize_update) - cum(kv_free.num_tokens) </div><div class="ln cont">+ cum(kv_get_computed_blocks.num_tokens)   (504/504)</div></pre>
<p class="pr">The KV cache conservation law, recovered from counts alone: prompt tokens admitted, plus detokenised outputs, minus every token freed.</p>
</div></div>'''
slides.append(Slide(
    "What drperf outputs", outputs_body,
    "derive gives an affine model, unexplained share and per-function attribution. learn gives relations between regions that held at every trigger.",
    ["An affine cost function per marked region", "Per-function attribution for every term", "Exact relations between independently marked states"],
    ["cost(superblocks) = 23.3*superblocks + 3,330.7", "attention.kv_tokens = admitted + generated - freed"],
))

method_formula = "pen_tensors       cost(num_reqs, max_len, total_tokens) = 3,299.1*num_reqs - 3*max_len + 309.9*total_tokens + 29,953.4      5.1%"
FORMULAS.append(method_formula)
method_body = '<div class="body case">'
method_body += block("annotation", "FIT", '<p class="pr">Declare caller-visible integer states, drive them independently, and reduce the unexplained share until the formula is a usable cost model.</p>' + mono([method_formula]))
method_body += block("surprise", "SURPRISE", '<p class="pr">Read the coefficient or constant that should not exist, then use its per-function attribution to locate the work.</p>' + mono(["309.9 per historical token  ->  PyArray_Pack / DiscoverDTypeAndShape / LONG_setitem"]))
method_body += block("optimise", "OPTIMISE", '<p class="pr">Change the named work, rerun the identical driver, and require the predicted coefficient to move while outputs stay identical.</p>' + mono(["309.9*total_tokens  ->  3.0*total_tokens"]))
method_body += '</div>'
slides.append(Slide(
    "The three-beat method: fit, surprise, optimise", method_body,
    "Measured means the third beat was run in an experiment copy or branch. Predicted means the first two beats exist and the expected coefficient is stated.",
    ["FIT — declare states and drive a reliable affine model", "SURPRISE — follow the unwanted term into its functions and source", "OPTIMISE — rerun before/after and watch that term move"],
    [method_formula, "309.9*total_tokens  ->  3.0*total_tokens"],
))


opt_groups = {
    "FTL": [
        ("no-clone locate()", "44.9*fill -> 0 at all three sites"),
        ("hash-keyed superblock map", "377k -> 304k per 32-block bind"),
    ],
    "vLLM": [
        ("alias-aware _prepare_inputs", "constant 464,928 -> 205,310"),
        ("cached request data", "4,916 -> 1,671 per running request"),
        ("array(\"L\") bounds check", "315.6 -> 109.2 per prompt token"),
        ("7-token detokenizer prime", "544.8 -> 227.5 per prompt token"),
        ("lazy request hashes", "254k -> 29k at 662 tokens; moved work"),
        ("input row memoryview", "275.2 -> 202.9 per prompt token"),
        ("metadata memo", "83,361 -> 44,595 per replacement"),
        ("batched condense", "47,366 -> 33,844 per moved row"),
        ("penalties tensor cache", "309.9 -> 3.0 per historical token"),
        ("skip early-return output call", "8,247.5 -> 6,989.4 per request-step"),
        ("prefix-scan memo", "22,800 -> 6,767 at 40 blocks"),
        ("remove hasher no-ops", "-993 each; -5.9% per block"),
        ("scheduler head memo", "23.7M -> 10.3M total, -56.5%"),
        ("one-pass preemption", "2,140.5 -> 422.8 per running request"),
        ("n-gram first-occurrence index", "22.9 -> 0.9 per context token"),
        ("set-based n-gram gather", "436,236 -> 221,638 at 64"),
        ("top-k only asking rows", "other-thread part -29.1%"),
    ],
    "WAN": [
        ("rotary table cache", "482,987 -> 58,612 per call"),
        ("prompt projection cache", "cond_embed -27.7%"),
        ("VAE list-then-cat", "927.5 -> 92.3 ms at production width"),
        ("lazy VAE posterior", "14.2 -> 0.349 per element"),
        ("callback locals", "14,019.5 -> 425 per name"),
        ("rotary apply layout", "5,764.3 -> 5,675.3 per token; wall -25%"),
        ("cross-attention K/V memo", "ax_qkv 378,712 -> 137,932"),
        ("short-pad T5 encode", "384.8M -> 8.07M at 8 tokens"),
    ],
}
cols = []
for group, items in opt_groups.items():
    lis = ''.join(f'<li><span class="what">{esc(name)}</span><span class="dash"> — </span><span class="res">{esc(result)}</span></li>' for name, result in items)
    cols.append(f'<div class="ogroup"><div class="gname">{group}</div><ul>{lis}</ul></div>')
opt_body = '<div class="body opt" style="display:grid;grid-template-columns:.82fr 1.22fr 1fr;gap:18px">' + ''.join(cols) + '</div>'
slides.append(Slide(
    "Optimisations that worked", opt_body,
    "Every line is a measured third beat from the current CASES.md. Lazy request hashes move work out of admission rather than removing it; rotary apply is the case where instructions and wall time disagree.",
    [f"{group} — {name}: {result}" for group, items in opt_groups.items() for name, result in items],
    [result for items in opt_groups.values() for _, result in items],
))


# FTL 1-8.
slides.extend([
case_slide(
    "FTL 1. The admission scan is not the 50 µs problem",
    "vllm-ditto/examples/cases/step_admission.py", 'with region("case_admission_limit", superblocks=resident):',
    "superblocks", "  cost(superblocks) = 23.3*superblocks + 3,330.7        [all superblocks]   blocks: 6 affine, 924 constant, 15 irregular (2.2% of cost)", "2.2%",
    "There is no bad term: 23.3 instructions per resident superblock makes the whole call about 6.2k at 120, while learn identifies resident exactly.",
    ["                23.3  <_ditto_ftl_core::directory::Directory>::stats  [_ditto_ftl_core.abi3.so]", "exact (96/96): case_admission_limit.superblocks = count(case_admission_limit) -2*cum(ftl_directory_invalidate.frees) +1"],
    "vllm-ditto/rust/ditto-ftl-core/src/directory.rs:876 — Directory::stats.",
    "Not needed: remove the proposed admission-cache fix from the list; under 14k instructions even at a few hundred superblocks.",
    notes="The useful result is a deletion from the optimisation list. The learned relation lets the cost be sized from commits and frees without another run.",
),
case_slide(
    "FTL 2. Three callers clone a block map to read a key or flags",
    "vllm-ditto/examples/cases/lookup_sites.py", 'with region("case_plan_load", fill=fill):',
    "fill", "  cost(fill) = 44.9*fill + -221.4        [all fill]   blocks: 55 affine, 134 constant, 0 irregular (0.0% of cost)", "0.0%",
    "Each lookup pays 44.9 instructions per resident io-block; caller slopes are 4×, 4× and 1×, exposing three clones.",
    ["                20.6  <alloc::collections::btree::map::BTreeMap<_, _, _> as core::clone::Clo  [_ditto_ftl_core.abi3.so]"],
    "vllm-ditto/rust/ditto-ftl-core/src/directory.rs:793 — superblock.clone().",
    "Measured: a no-clone locate() makes every fill coefficient exactly zero; digest c4529c2f9fbf6a75 unchanged.",
    ["derive case_plan_load            cost(fill) = 167.6*fill + 3,025.1", "derive case_plan_load            cost(fill) = 0*fill + 877.3        blocks: 0 affine, 1960 constant, 51 irregular (77.9% of cost)"],
    "The formula sizes the full-context saving at about 140k instructions, an order of magnitude below the old estimate, and identifies two additional call sites.",
),
case_slide(
    "FTL 3. Small-scale bind_batch numbers favour the loser",
    "vllm-ditto/rust/ditto-ftl-core/src/directory.rs", '"ftl_directory_bind_batch",',
    "requests", "  cost(requests) = 2,237.6*requests + 3,877.5        [all requests]", "19.5% after four driver repairs",
    "The old core scales with the whole directory; the current one spends about 44% of a bind in string-keyed B-tree search, whose leaf occupancy is internal state.",
    ["            15,606.8  __memcmp_avx2_movbe  [libc.so.6]"],
    "vllm-ditto/rust/ditto-ftl-core/src/directory.rs:152 — the string-keyed search under apply_batch.",
    "Measured: change only the map keying; the directory-size climb disappears and mean cost falls 377,313 → 304,471 per 32-block bind.",
    ["the 1,371,318 / 377,313 / 304,471 instructions per bind are measured"],
    "At five superblocks the pre-rewrite clone can look cheaper; at production it is eight times worse. The fit itself needed repeated exact points and one code path.",
),
case_slide(
    "FTL 4. count() round-trips every counter",
    "vllm-ditto/examples/cases/count_stats.py", 'with region("case_count", entries=entries):',
    "entries", "  cost(entries) = 1,896.6*entries + 1,060.1        [all entries]   blocks: 359 affine, 700 constant, 47 irregular (22.8% of cost)", "22.8%",
    "A single counter bump reads and rewrites the whole stats dict: about 58k instructions with roughly 30 counters.",
    ["               159.5  PyDict_SetItem  [python3.12]"],
    "vllm-ditto/rust/ditto-ftl-core/src/lib.rs:1016-1021 — PyFTLController::pull_stats / sync_stats.",
    "Predicted: keep counters in Rust and materialise the Python dict only on read, removing about 55k per count().",
    notes="The formulas reverse the old priority order: count first, plan_load second, admission not at all.",
),
case_slide(
    "FTL 5. Container size is not container state",
    "examples/hashmap_internal.py", 'with region("hm_lookup", n=n, size=size):',
    "n, size", "hm_insert   cost(n, size) = 1,970.2*n + 0*size + 2,624      10.6% irregular", "10.6%",
    "size has coefficient zero; the residue is a resize step near two-thirds load, not a slope in entries.",
    ["    irregular by function: PyDict_Contains 2,020 | PyDict_SetItem 1,293 | _Py_HashBytes 672"],
    "examples/hashmap_internal.py:41 — the insert region; CPython's table headroom is hidden state.",
    "No code change: widen the sweep. The n term holds while the proxy rehash coefficient moves from 5 to 1.7.",
    ["sizes 40..340,  n 8..32:   cost = 2,014.7*n + 5*rehash   + 2,518    6.6% irregular", "sizes 20..1500, n 4..64:   cost = 1,957.8*n + 1.7*rehash + 1,428   10.5% irregular"],
    "This forty-line analogue explains why B-tree node fill and hash-table doubling belong in the trace when callers cannot declare them.",
),
case_slide(
    "FTL 6. _ordered_blocks hides a quadratic clone",
    "vllm-ditto/examples/cases/ordered_blocks.py", 'with region("case_ordered_blocks", nblocks=n, square=n * n):',
    "nblocks, square=nblocks²", "case_ordered_blocks   cost(nblocks, square) = 2,049.7*nblocks + 416.9*square + 4,942.8    (16.1% irregular)", "16.1%",
    "The blocks getter clones and converts the Rust BTreeMap once per offset: 417 instructions per offset×entry for a tuple.",
    ["               127.4  PyDict_SetItem", "                45.2  <btree::map::IntoIter<usize, i64>>::dying_next"],
    "vllm-ditto/src/integration/vllm/compaction.py:231 via the pyo3 getter at lib.rs:682.",
    "Predicted from the one-getter region: the square coefficient becomes zero; 497k → 54k per call at 32 blocks.",
    ["case_ordered_once     cost(nblocks, square) = 1,499.7*nblocks + -0.196*square + 6,363.1   (1.3% irregular)"],
    "A flat production state can still be the state that squares. This was invisible because callers opened their regions after the getter call.",
),
case_slide(
    "FTL 7. ensure_row gets slower as the pool fills",
    "vllm-ditto/examples/cases/ensure_row_slabs.py", '"case_ensure_row", full_slabs=row // slots, new_slab=int(row % slots == 0)',
    "full_slabs, new_slab", "case_ensure_row   cost(full_slabs, new_slab) = 35.1*full_slabs + 381.2*new_slab + 2,397.2   (35.1% irregular, 3 blocks, 272..4,442 per call)", "35.1%",
    "Every stored row walks the prefix of exhausted slabs; the regression arrives with occupancy, not workload size.",
    ["                33.8  <ManagerCore>::alloc_row_slot"],
    "vllm-ditto/rust/ditto-ftl-core/src/handle_manager.rs:453 — alloc_row_slot.",
    "Predicted: keep only non-full slabs and make the 35.1 coefficient zero; about 540k saved per 32-row store at a full 480-slab pool.",
    notes="The residue is bounded slab-index growth. A future driver should free and refill to measure find_run over a fragmented free set.",
),
case_slide(
    "FTL 8. _clear_digests scans the table once per retired block",
    "vllm-ditto/examples/cases/digest_scan.py", 'with region("case_drop_key", digests=digests):',
    "digests", "case_drop_key   cost(digests) = 288.4*digests + 7,395.3        [all digests]   0.0% irregular", "0.0%",
    "The whole digest table is copied and compared per victim: about 1.1M per block and 35M per retired superblock at production size.",
    ["                 211  _PyEval_EvalFrameDefault", "                  47  PyDict_Clear"],
    "vllm-ditto/src/integration/vllm/compaction.py:253-258 — _clear_digests from _drop_key.",
    "Predicted: maintain a key-to-identities index and make the digests coefficient zero.",
    notes="The outer invalidate region originally derived 0*digests because its nested region did not declare digests; this becomes the first tool caveat later.",
),
])


# vLLM A-H2, with two slides each for A and B.
slides.extend([
case_slide(
    "vLLM A1. _prepare_inputs pays a fixed 465k tax per step",
    "vllm-cpu-src/vllm/v1/worker/gpu_model_runner.py", 'with perfmark.region("prepare_inputs", num_reqs=len(scheduler_output.num_scheduled_tokens), num_tokens=scheduler_output.total_num_scheduled_tokens):',
    "num_reqs, num_tokens", "prepare_inputs        cost(num_reqs, num_tokens) = 758.2*num_reqs + 106.9*num_tokens + 464,927.7                    20.9% irregular", "20.9%",
    "The fixed tax is ten times the request-dependent work at 64 requests; on CPU, copy_to_gpu dispatches copies from tensors onto themselves.",
    ["constant: torch dispatch; __tls_get_addr 13.2k; pthread_mutex_lock/unlock 17.5k"],
    "vllm-cpu-src/vllm/v1/utils.py:139 — copy_to_gpu; gpu_model_runner.py:2037-2262 — the small tensor operations.",
    "Measured: guard the aliasing CPU path and reuse numpy values; constant 464,927.7 → 205,310.3 (-56%), outputs identical.",
    ["prepare_inputs   before  758.2*num_reqs + 106.9*num_tokens + 464,927.7     20.9% irregular", "                 after   721.3*num_reqs +  55.2*num_tokens + 205,310.3     26.0%"],
    "The matching 256k drop in execute_model is the composition identity: the fixed work was truly removed, not reassigned.",
),
case_slide(
    "vLLM A2. CachedRequestData rebuilds unchanged state every step",
    "vllm-cpu-rb/vllm/v1/core/sched/scheduler.py", 'with perfmark.region("cached_request_data", num_running=len(running_reqs), num_resumed=len(resumed_reqs)):',
    "num_running, num_resumed", "cached_request_data   cost(num_running, num_resumed) = 4,907.1*num_running + 0*num_resumed + 13,107.9                  3.8%", "3.8%",
    "A fresh tuple of lists, two generators and the entire CachedRequestData are rebuilt for every running request even when only one integer changed.",
    ["per running request: PyFunction_New 242; interpreter cost 2,826"],
    "vllm-cpu-rb/vllm/v1/core/sched/scheduler.py:1525; kv_cache_manager.py:77-92 — get_block_ids.",
    "Measured: skip get_block_ids on the fifteen decode steps in sixteen with no new blocks and hoist loop invariants; 4,916.2 → 1,671 per request.",
    ["cached_request_data   4,916.2*num_running + 13,027.5    4.2%   ->   1,671*num_running + 15,247.4    3.7%"],
    "At 64 running requests, 328k → 122k per step; the consumer formula remains unchanged.",
),
case_slide(
    "vLLM B1. Admission scans every prompt twice, then decodes it again",
    "vllm-cpu-req/vllm/v1/engine/input_processor.py", 'with perfmark.region("process_inputs", num_tokens=(len(prompt.get("prompt_token_ids") or ()) if isinstance(prompt, dict) else (len(prompt) if isinstance(prompt, str) else 0))):',
    "num_tokens", "process_inputs       cost(num_tokens)   = 315.6*num_tokens + 372,902.8       [num_tokens >= 62]       4.3%", "4.3%",
    "The token slope is Python max/min over IDs; engine_add_request then primes a detokenizer from the entire prompt.",
    ["    per-token: PyObject_RichCompareBool 116, PyLong_AsSsize_t 74, PyList_Size 33, PyIter_Next 30"],
    "vllm-cpu-req/vllm/v1/engine/input_processor.py:485-486; detokenizer.py:184.",
    "Measured: array(\"L\") makes one C pass and a decoder-specific seven-token prime removes the second slope; outputs identical.",
    ["process_inputs       315.6*num_tokens + 372,902.8   ->   109.2*num_tokens + 386,776.3", "engine_add_request   544.8*num_tokens + 603,204.4   ->   227.5*num_tokens + 618,328.6"],
    "np.asarray and all(map(...)) lost in measurement; only array(\"L\") beat the status quo.",
),
case_slide(
    "vLLM B2. Request.__init__ is a 128-token staircase",
    "vllm-cpu-req/vllm/v1/request.py", 'with perfmark.region("request_init", num_prompt=(len(prompt_token_ids) if prompt_token_ids is not None else 0), num_blocks=(len(prompt_token_ids) // getattr(block_hasher, "hash_block_size", 0) if prompt_token_ids and getattr(block_hasher, "hash_block_size", 0) else 0)):',
    "num_prompt, num_blocks", "request_init   cost(num_prompt, num_blocks) = -4.0*num_prompt + 43,273.3*num_blocks + 25,363.4     57.7% -> 12.7% irregular", "12.7% after adding num_blocks",
    "The apparent 79-per-token slope is really 43,273 per 128-token block: eager block hashes plus a required prompt copy.",
    ["    per block: _PyObject_GC_Resize 19,526, _PyEval_EvalFrameDefault 3,714, PyLong_AsLongAndOverflow 2,706, PyList_AsTuple 1,635, SHA1_Init 1,240"],
    "vllm-cpu-req/vllm/v1/request.py:211 — update_block_hashes().",
    "Measured: make block_hashes lazy. At 662 tokens Request.__init__ falls 254,024 → 29,128; the hash work moves into schedule and total work is unchanged.",
    ["request_init         78.5*n + 23,998 (57.7%)     ->   -2.5*n + 1,086*blocks + 18,829 (34.1%)     662 tokens: 254,024 -> 29,128", "engine_add_request   227.5*n + 618,329 (13.6%)   ->   136.8*n + 620,710 (7.0%)                   662 tokens: 1,011,290 -> 760,785"],
    "This improves time-to-first-schedule, not total instructions: schedule's increase matches admission's decrease.",
),
case_slide(
    "vLLM C. Input-batch churn bills arrivals per token and departures per row",
    "vllm-cpu-batch/vllm/v1/worker/gpu_input_batch.py", 'with perfmark.region("add_request", num_prompt=(len(request.prompt_token_ids) if request.prompt_token_ids is not None else 0), num_blocks=sum(len(b) for b in request.block_ids)):',
    "num_prompt, num_blocks", "add_request       cost(num_prompt, num_blocks) = 275.2*num_prompt + 0*num_blocks + 46,456.8                                      2.8%", "2.8%",
    "Admission converts a Python list element by element into int32; condense pays for roughly seventeen numpy dispatches per moved row, not bytes copied.",
    ["    per-prompt-token: PyArray_DiscoverDTypeAndShape_Recursive 85, PyArray_Pack 69, INT_setitem 58, PyLong_AsLong 22"],
    "vllm-cpu-batch/vllm/v1/worker/gpu_input_batch.py:376/382/387 and :765-786.",
    "Measured: cached row memoryviews, scalar metadata memoisation, and batched moves only above the measured 2.7-row crossover.",
    ["add_request        275.2*num_prompt + 46,457    ->   202.9*num_prompt + 42,629       (long prompts: 332.7 -> 269.6)", "refresh_metadata   at the steady replace (8 requests, 1 added): 83,361 -> 44,595 per call   (-46.5%)", "condense           per moved row, bursts of >= 3: 47,366 -> 33,844   (-28.5%); single departures unchanged"],
    "Ten plausible conversion candidates were measured; most were regressions.",
),
case_slide(
    "vLLM D. Penalties rebuild the whole generated history every step",
    "vllm-cpu-samp/vllm/v1/sample/ops/penalties.py", 'with perfmark.region("pen_tensors", num_reqs=len(output_token_ids), max_len=max((len(o) for o in output_token_ids), default=0), total_tokens=sum(len(o) for o in output_token_ids)):',
    "num_reqs, max_len, total_tokens", "pen_tensors       cost(num_reqs, max_len, total_tokens) = 3,299.1*num_reqs - 3*max_len + 309.9*total_tokens + 29,953.4      5.1%", "5.1%",
    "Every step re-encodes every token generated so far: 309.9 instructions per historical token, hence a quadratic generation cost.",
    ["    per-total_tokens: PyArray_Pack 69, PyArray_DiscoverDTypeAndShape_Recursive 64, LONG_setitem 57"],
    "vllm-cpu-samp/vllm/v1/sample/ops/penalties.py:47 into torch_utils.py:731.",
    "Measured: cache one int64 array per request and extend it; 309.9 → 3.0 per token, identical outputs, about 10G saved at 64×1,000.",
    ["pen_tensors       before  3,299.1*num_reqs - 3*max_len + 309.9*total_tokens + 29,953.4     5.1% irregular", "                  after     376.8*num_reqs + 0*max_len +   3.0*total_tokens + 29,213.6    66.7%"],
    "The high after-percentage is the remaining cache branch inside a region now one tenth the size.",
),
case_slide(
    "vLLM E. The output-path surprise was mostly its markers",
    "vllm-cpu-out/vllm/v1/engine/output_processor.py", 'with perfmark.region("process_outputs", num_outputs=len(engine_core_outputs), num_active=len(self.request_states)):',
    "num_outputs, num_active", "out/vllm_out_clean  (2 markers)   process_outputs = 30,153.5*num_outputs + 13,663.7*num_active + 2,952.9     26.2%", "26.2% with two markers; 52.9% with seven",
    "Five nested Python markers contributed roughly 11k construction instructions each to the parent, creating a plausible but false per-request cost.",
    ["_split_state -> isinstance(v, numbers.Integral); only 582 in-region instructions are calibrated away"],
    "/home/ubuntu/drperf/perfmark/python/perfmark.py:124-132; output_processor.py:293-296 already returns early.",
    "Measured after removing the measurement artifact: skip the early-return call; true FINAL_ONLY cost 8,247.5 → 6,989.4 (-15.3%), 752/752 records identical.",
    ["8,247.5 to", "6,989.4 (-15.3%), about 80k per step at 64 requests"],
    "The negative constant and 53% residue were the warnings. A lazy text join was also measured and rejected, as the original formula predicted.",
),
case_slide(
    "vLLM F. Prefix caching spends more on no-op calls than hashing",
    "vllm-cpu-kv/vllm/v1/core/block_pool.py", 'with perfmark.region("kv_cache_full_blocks", num_new=num_full_blocks - num_cached_blocks, num_cached=num_cached_blocks):',
    "num_new, num_cached", "kv_cache_full_blocks    cost(num_new, num_cached) = 7,267.2*num_new + 0*num_cached + 4,521.4                          7.1% irregular", "7.1%",
    "Fifteen calls in sixteen hash nothing; a separate per-step scan walks the shared prefix. The real block cost is pickle+SHA, not the three copies.",
    ["    per new block: _PyEval_EvalFrameDefault 3,992, _Py_HashBytes 198, PyBytes_FromString 160, PyLong_AsUnsignedLong 153"],
    "vllm-cpu-kv/vllm/v1/core/kv_cache_utils.py:695 and hashing.py:39; single_type_kv_cache_manager.py:823-830.",
    "Measured: remove no-op hasher calls (-993 each, -5.9% per block) and memoise the prefix scan, 22,800 → 6,767 at 40 blocks.",
    ["kv_common_prefix_blocks   0*num_running + 0*num_blocks + 1,765.4    70.0%   ->   0*num_running + 0*num_blocks + 2,934.3    8.8%", "kv_cache_full_blocks               7,266*num_new + 4,480        6,839.6*num_new + 4,765.2   (-5.9% per block)"],
    "The conservation relation is exact at 504/504 attention calls: resident KV = admitted + generated - freed.",
),
case_slide(
    "vLLM G. Scheduler pressure: waiting is free; re-evaluation and preemption are not",
    "vllm-cpu-sched/vllm/v1/core/sched/scheduler.py", 'with perfmark.region("sched_preempt", running=len(self.running), preempted=len(preempted_reqs), num_new_tokens=num_new_tokens):',
    "running, preempted, num_new_tokens", "sched_preempt         cost(running, preempted, num_new_tokens) = 2,140.5*running + 0 + 0 + 94,086.3   14.9%", "14.9%",
    "The earlier waiting slope was aliasing: under max_num_seqs saturation it is exactly zero. Real costs are a ~200k head retry and five scans of running per victim.",
    ["           per running: _PyTuple_Resize 358.6, PyTuple_GetSlice 210.6, PyObject_RichCompareBool 203.8, PyTuple_New 139.4"],
    "vllm-cpu-sched/vllm/v1/core/sched/scheduler.py:648-664 and :762-771.",
    "Measured with identical admission order, victims and peek counts: head retries -56.5%; preemption slope 2,140.5 → 422.8.",
    ["sched_preempt       2,140.5*running + 94,086.3    14.9%   ->   422.8*running + 98,683.2    9.9%", "sched_waiting_eval  triggers 168 -> 59 (36 admissions + 23 forced by real frees); total 23,714,077 -> 10,319,210 (-56.5%)"],
    "At 64 running with KV full: about 123k saved per step by the memo plus about 110k per preemption.",
),
case_slide(
    "vLLM H. N-gram speculation cannot pay on this CPU backend",
    "vllm-cpu-spec/vllm/v1/spec_decode/ngram_proposer.py", 'with perfmark.region("ngram_scan", num_valid=num_ngram_requests,',
    "num_valid, ctx, k, tag", "ngram_scan         cost(num_valid, ctx, k, tag) = 9,926.9*num_valid + 18*ctx + 0*k + 24,536.1        6.6% irregular", "6.6%",
    "The proposer rescans the entire context; meanwhile execute_model saves at most 3.2% per scored token and rejection sampling costs 61% more per output token.",
    ["                       per ctx: 25.1 <no module> (the numba KMP kernel)"],
    "vllm-cpu-spec/vllm/v1/spec_decode/ngram_proposer.py:263-296 — reversed-context KMP scan.",
    "Measured: K=3 and K=5 are a throughput wash; a first-occurrence index makes the context slope 22.9 → 0.9 and set membership removes gather's n² term.",
    ["ngram_scan   before  9,856*num_valid + 22.9*ctx + 32,441   (1.4%)", "             after   11,909*num_valid + 0.9*ctx + 34,705   (3.3%), plus a build of ~300 per token indexed, once per request"],
    "At 64×1,000, scan 2.13M → 0.88M and gather 436k → 222k per step; the 67.6G forward still dominates.",
),
case_slide(
    "vLLM H2. One logprobs request bills every row",
    "mark_spec.py", '("vllm/v1/sample/sampler.py", "Sampler", "gather_logprobs", "gather_logprobs",',
    "num_reqs, num_logprobs, tag", "gather_logprobs   cost(num_reqs, num_logprobs, tag) = 1,009,495.2*num_reqs + 115,043.4           57.2%", "57.2%",
    "Stop strings are flat at 4,196.8 instructions and min_tokens is free when unused; the real surprise is top-k over all rows when only one asks.",
    ["                      per num_reqs: 913,749.3 at::native::AVX2::topk_impl_loop, 88,004 kernel._omp_fn.0"],
    "vllm-cpu-spec/vllm/v1/sample/sampler.py:337 — torch.topk(logprobs, ..., dim=-1). Marker shown is the exact stop-set injection spec.",
    "Measured: top-k only the asking rows; other-thread work -29.1% with byte-identical top-5 dictionaries, but wall time is unchanged at eight threads.",
    ["steady-state row       total/call     64,131,276             59,409,231     -7.4%", "  of which other-thr   16,205,197                            11,489,509    -29.1%"],
    "The result states when it matters: at 64 requests with eight asking, 56×1.18M instructions disappear per step.",
),
])


# Wan main section and B/C/D.
slides.extend([
case_slide(
    "Wan 2.1. Fixed host work repeats every denoising step",
    "videogen/wan-fit/diffusers/models/transformers/transformer_wan.py", 'with perfmark.region("rope", frames=hidden_states.shape[2], height=hidden_states.shape[3], width=hidden_states.shape[4]):',
    "frames, height, width", "rope                = 29.9*frames + 0*height + 0*width + 312,303.4                                             35.2%", "35.2%",
    "The rotary table is rebuilt every forward; the fixed prompt is re-projected every step and CFG branch; VAE output concatenation copies quadratically in frames.",
    ["WanRotaryPosEmbed.forward: 312k fixed", "text_embedder: cond_embed 617k fixed", "VAE torch.cat: somatcopy irregular work"],
    "videogen/wan-fit/diffusers/models/transformers/transformer_wan.py:395/668; pipeline_wan.py:347; autoencoder_kl_wan.py:1197-1205.",
    "Measured with byte-identical latents/video: rotary 482,987 → 58,612 (-87.9%), cond_embed -27.7%, and list-then-cat is 10× faster at production width.",
    ["             mean per call 482,987 -> 58,612 (-87.9%); first call 782k, every cached call 13,679", "             mean per call 780,728 -> 564,337 (-27.7%); the residue is the timestep embedding, which does change per step"],
    "The VAE rising per-frame total was mostly first-chunk amortisation, not cat; a production-width micro-measurement kept the cat fix on different evidence.",
),
case_slide(
    "Wan B. Unmarked paths: rotary apply, callbacks and the VAE posterior",
    "videogen/wan-more/diffusers/models/transformers/transformer_wan.py", 'with perfmark.region("attn_rope", tag=902, seq_len=hidden_states.shape[1], heads=attn.heads, batch=hidden_states.shape[0]):',
    "seq_len, heads, batch", "attn_rope      = 5,764.3*seq_len + 0*heads + 0*batch + 320,506.1        0.4% irregular", "0.4%",
    "Rotary application spends 45 instructions per rotated element in iterator overhead; callbacks call locals() per name; the VAE computes exp for posterior values mode() never reads.",
    ["    per token: c10::function_ref callback 3,064, serial_for_each 1,264, callback 896, DimCounter::increment 352"],
    "videogen/wan-more/diffusers/models/transformers/transformer_wan.py:104-118; pipeline_wan.py:640; vae.py:693-694.",
    "Measured, outputs identical: posterior slope 40× lower, callback per-name 33× lower, rotary instructions -1.5% but wall -25%; blend loops deliberately unchanged.",
    ["dgd_init        14.2*numel + 114,569.1      2.7%   ->   0.349*numel + 70,030      1.8%     (40x on the slope; 9.31 -> 0.10 ms per production encode)", "callback_step   14,019.5*ninputs + 10,872.6 5.4%   ->   425*ninputs + 26,940      0.0%     (33x per name)", "attn_rope       5,764.3*seq_len + 320,506   0.4%   ->   5,675.3*seq_len + 323,964 0.4%     (-1.5% instructions; 210.9 -> 157.7 ms per q or k at production, -25% wall)"],
    "The rotary result is the deck's explicit reminder that instructions are not elapsed time.",
),
case_slide(
    "Wan C. Cross-attention re-projects the same prompt 3,000 times",
    "videogen/wan-tax/diffusers/models/transformers/transformer_wan.py", 'with perfmark.region("attn", batch=hidden_states.shape[0], seq_len=hidden_states.shape[1], heads=attn.heads, kv_len=(hidden_states.shape[1] if encoder_hidden_states is None else encoder_hidden_states.shape[1])):',
    "batch, seq_len, heads, kv_len", "block               = 0*batch + 12,424.7*seq_len + 0*text_len + 2,172,643.1                                    79.5%", "79.5% in the parent before 55 temporary probes",
    "The per-layer constant contains ax_qkv at 378,712, including 130,563 sgemm instructions: fixed prompt K/V projected 30 layers × 100 forwards instead of 60 times.",
    ["ax_qkv: mkl_blas_def_sgemm_kernel_0_zen 130,563 of 378,712"],
    "videogen/wan-tax/diffusers/models/transformers/transformer_wan.py:45 — attn.to_k / attn.to_v on encoder_hidden_states.",
    "Measured with identical latents/video: memoise K/V by tensor identity and version; ax_qkv -63.6%, host tax -9.8% and 14.2 TFLOP removed per production video.",
    ["ax_qkv (cross-attn K/V)   378,712 -> 137,932   (-63.6%)", "wan_call             ... + 10,750,513.4*steps + 27,537,440   ->   ... + 9,502,853.1*steps + 27,488,895"],
    "The 55 probes also removed repeated timestep/modulation work, fp32 upcasts, Dropout(p=0), rotary slices and identity type_as calls.",
),
case_slide(
    "Wan D. The text encoder pays for padding, quadratically",
    "videogen/wan-t5/diffusers/pipelines/wan/pipeline_wan.py", 'with perfmark.region("t5_enc", tag=926, ptok=_b * max_sequence_length, sq=_b * max_sequence_length * max_sequence_length // 64, lsq=max_sequence_length * max_sequence_length // 64):',
    "rtok, ptok, sq, lsq", "t5_enc      24,980*ptok + 46,060.5*sq + 27,714.6*lsq + 6,807,980                      17.7%             (with lsq)", "17.7%",
    "The real prompt coefficient is -1.2: effectively free. Work follows padded tokens, padded pairs and an L² relative-position bias rebuilt per layer.",
    ["    ptok: mkl sgemm_kernel 10,497, mkl_vml sTanh 7,232        (projections, FFN)", "    sq:   mkl sgemm_pst 17,664                                (attention, per padded token pair)", "    lsq:  index_select_out_cpu 4,608, mkl_vml sLn 3,320       (the relative-position bias, rebuilt as (1, heads, L, L) in every layer)"],
    "videogen/wan-t5/diffusers/pipelines/wan/pipeline_wan.py:199; hf/transformers/src/transformers/models/umt5/modeling_umt5.py:347 — position-bias construction.",
    "Measured: encode at the longest real length rounded to eight (floor 16), then zero-pad the embeddings; 384.8M → 8.07M at eight real tokens, bit-identical.",
    ["t5_embeds   25,688.3*ptok + 46,070.5*sq + 27,692*lsq + 6,565,689   ->   693.5*ptok + 34.8*sq - 27.7*lsq + 6,694,193", "per call, batch 1, max_len 512:  8 real tokens 384.8M -> 8.07M (47.7x); 32 -> 8.46M (45.5x); 128 -> 17.3M (22.2x); 480 -> 124.5M (3.1x)"],
    "Cross-attention keys cannot be shortened without changing the model: it was trained attending to its padding.",
),
])


# Fit pass.
fit_rows = [
    ("Wan tiled_encode_r", "100.0", "0.1", "tiles, area_frames; stride-aligned grid, frames=1"),
    ("Wan tiled_decode_r", "100.0", "0.4", "tiles, area_frames"),
    ("Wan pp_denorm", "96.0", "0.0", "pixels_k"),
    ("Wan prompt_clean_r", "93.7", "3.8", "chars, nonascii"),
    ("Wan postprocess_frames", "90.7", "0.1", "pixels_k"),
    ("vLLM detok_update", "53.0", "0.0", "is_first, have_k"),
    ("Wan attn_rope", "40.9", "0.0", "sh=seq_len*heads, head_dim"),
    ("vLLM kv_allocate_slots", "49.0 / 21.3", "18.2 / 1.4", "new_blocks, is_prefill"),
    ("Wan postprocess_video", "88.0", "16.2", "pixels_k"),
    ("Wan vae_decode_chunk / frame_loop", "99.3 / 99.4", "4.3 / 4.3", "first, area_k, area_frames_k"),
    ("Wan vae_decode_frames / stage", "99.7 / 99.7", "7.4 / 7.4", "chunks, area_k, area_frames_k"),
    ("Wan vae_encode / encode_chunk", "99.2 / 97.2", "2.1 / 2.0", "first, area_k, area_frames_k"),
    ("Wan wan_call", "98.7", "11.5", "steps, sseq, ssq_k, pix_k"),
]
rows_html = ''.join(f'<tr><td>{esc(r)}</td><td>{b}</td><td>{a}</td><td><code style="font-size:13px">{esc(st)}</code></td></tr>' for r,b,a,st in fit_rows)
fit_body = '<div class="body tblwrap"><table class="sum"><colgroup><col style="width:29%"><col style="width:12%"><col style="width:12%"><col style="width:47%"></colgroup><thead><tr><th>region</th><th>before %</th><th>after %</th><th>states added</th></tr></thead><tbody>' + rows_html + '</tbody></table></div>'
slides.append(Slide(
    "Fit pass: regions taken from above 20% to under it", fit_body,
    "Only rows whose after measurement is under 20% are shown. update_from_output improved in one regime but remained 34.8% in saturation, so it is omitted from this success table.",
    [f"{r}: {b}% → {a}% with {st}" for r,b,a,st in fit_rows],
    ["before unexplained %  ->  after unexplained %"],
))

resist_body = '<div class="body rules"><p class="intro">Four structural limits remained after re-declaration and new grids.</p><ol>'
resist = [
    ("schedule: five code paths", "No single plane fits admission, saturation, chunking, preemption and steady decode. Trace fit reaches R² 0.993, but the actionable sub-regions fit at 1.2-16.5%."),
    ("pp_stack: alternating memcpy blocks", "Two __memcpy_avx_unaligned_erms blocks take turns with numpy placement. Each is irregular; their sum is affine to 0.03%."),
    ("cpu_flash_attention: three template instantiations", "<32,512>, <64,512> and <256,512> are selected by seq_len thresholds. Restricting to one band takes attn_sdpa to 15.2%."),
    ("Four-state limit", "KEY_STATES=4 forces trade-offs: process_outputs needs two sizes, first-token and output-kind; wan_call had to drop chunks."),
]
for h, b in resist:
    resist_body += f'<li><span class="rhead">{esc(h)}.</span> <span class="rbody">{esc(b)}</span></li>'
resist_body += '</ol>' + mono(["schedule: five paths  |  pp_stack: two alternating blocks  |  flash attention: three seq_len bands  |  KEY_STATES=4"]) + '</div>'
slides.append(Slide(
    "Fit pass: what would not become one affine model", resist_body,
    "The remedy is segmentation or a tool change, not inventing another correlated state. process_outputs also remains at 25% because the necessary state signature exceeds four integers.",
    [f"{h}: {b}" for h,b in resist],
    ["<32,512>  <64,512>  <256,512>  |  KEY_STATES=4"],
))


# Driver rules and required caveats.
driver_rules = [
    ("Repeat every state point", "Singleton points have no stable mean; allocator jitter becomes unexplained work."),
    ("Use exact sizes, not buckets", "A bucket spans real sizes and turns within-point spread into variation."),
    ("Keep one code path per region", "Fixture creation and the measured operation cannot share one plane."),
    ("Drop confounded states", "A state moving with batch size may explain the driver rather than the code."),
    ("Declare computed states", "square=n*n exposed case 6; products are valid when the code really scales with them."),
]
driver_body = '<div class="body rules"><p class="intro">Seven declarations did not move bind_commit. Four driver repairs took it from 62-70% unexplained to 19.5%.</p><ol>'
for h,b in driver_rules:
    driver_body += f'<li><span class="rhead">{esc(h)}.</span> <span class="rbody">{esc(b)}</span></li>'
driver_body += '</ol>' + mono(["  cost(requests) = 2,237.6*requests + 3,877.5        [all requests]", "blocks: 192 affine, 648 constant, 33 irregular (19.5% of cost)"]) + '</div>'
FORMULAS.append("  cost(requests) = 2,237.6*requests + 3,877.5        [all requests]")
slides.append(Slide(
    "Driver rules: fix the experiment before the declaration", driver_body,
    "A high unexplained share is first evidence about how the region was exercised, and only then about the code.",
    [f"{h}: {b}" for h,b in driver_rules],
    ["cost(requests) = 2,237.6*requests + 3,877.5  |  19.5% unexplained"],
))

caveats = [
    ("Nested regions average away outer-state dependence", "Nested means are keyed by the inner region's states. Declare the outer state inside, or measure the inner work separately."),
    ("About 11k marker-construction instructions land in the parent", "A four-state Python region spends ~3.8 µs in _split_state; calibration removes only the 582 instructions inside the child."),
    ("Slot-table overflow silently misattributes", "vLLM's ~556k blocks overflow the 131,072-slot default and collapse into the last slot. Use --max-slots and reject warned runs."),
    ("The key cache compares state names by pointer", "Identical state signatures can merge regions. Give signatures unique values until client/drperf.c compares string contents."),
    ("derive drops small states beside a huge computed basis", "A sq column around 6.4e12 makes steps look dependent. Scale products, for example sq_k = seq_len*kv_len // 1024."),
    ("Late attach can land inside a JIT", "NgramProposer construction triggered a 6.06B-instruction numba compile. Disable markers through construction and warm-up."),
]
caveat_body = '<div class="body rules"><p class="intro">Each caveat changed a reading before it changed a decision.</p><ol>'
for h,b in caveats:
    caveat_body += f'<li><span class="rhead">{esc(h)}.</span> <span class="rbody">{esc(b)}</span></li>'
caveat_body += '</ol></div>'
slides.append(Slide(
    "Tool caveats: six ways a plausible formula can mislead", caveat_body,
    "Additional notes in CASES.md cover marker-free predict comparisons, fast-path mis-pairing, periodic states, the 128-point cap, and OpenMP own-thread spin.",
    [f"{h}: {b}" for h,b in caveats],
    ["--max-slots 2097152  |  sq_k = seq_len*kv_len // 1024  |  KEY_STATES=4"],
))


def inline_md(s: str) -> str:
    parts = re.split(r'(`[^`]*`)', s)
    out = []
    for p in parts:
        if p.startswith('`') and p.endswith('`'):
            out.append(f'<code style="font-size:13px">{esc(p[1:-1])}</code>')
        else:
            out.append(esc(p))
    return ''.join(out)


case_lines = CASES.read_text().splitlines()
table_start = next(i for i, line in enumerate(case_lines) if line == "| case | fit | surprise | optimise |")
summary_rows: list[list[str]] = []
for line in case_lines[table_start + 2:]:
    if not line.startswith("|"):
        break
    cells = [c.strip() for c in line.strip().strip("|").split("|")]
    if len(cells) != 4:
        raise RuntimeError(f"summary row did not have four cells: {line}")
    summary_rows.append(cells)

# Four pages keep every formula at >=13px and every cell readable.
chunks = [summary_rows[:5], summary_rows[5:10], summary_rows[10:15], summary_rows[15:]]
for index, chunk in enumerate(chunks, 1):
    rows = ''.join('<tr>' + ''.join(f'<td>{inline_md(cell)}</td>' for cell in row) + '</tr>' for row in chunk)
    body = '<div class="body tblwrap"><table class="sum"><colgroup><col class="c1"><col class="c2"><col class="c3"><col class="c4"></colgroup><thead><tr><th>case</th><th>fit</th><th>surprise</th><th>optimise</th></tr></thead><tbody>' + rows + '</tbody></table></div>'
    slides.append(Slide(
        f"The cases as slides: fit, surprise, optimise ({index}/4)", body,
        "Cells are rendered verbatim from the summary table at the top of CASES.md. Measured means the third beat was run; predicted means only the first two beats exist.",
        [" | ".join(row) for row in chunk],
        [row[1] for row in chunk],
    ))


def wrap_section(slide: Slide, first: bool = False) -> str:
    if first:
        core = slide.body
    else:
        core = f'<h2>{esc(slide.title)}</h2>\n{slide.body}'
    return f'<section class="slide">{core}<aside class="notes">{esc(slide.notes)}</aside></section>'


def build_html() -> None:
    old = HTML_OUT.read_text()
    prefix_end = old.index('<div id="stage">') + len('<div id="stage">')
    suffix_start = old.rindex('</div>\n<div id="counter">')
    prefix = old[:prefix_end]
    suffix = old[suffix_start:]
    content = "\n" + "\n".join(wrap_section(s, i == 0) for i, s in enumerate(slides)) + "\n"
    HTML_OUT.write_text(prefix + content + suffix)


def build_pptx() -> bool:
    try:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
        from pptx.util import Inches, Pt
    except Exception:
        return False

    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    ink = RGBColor(20, 22, 26)
    dim = RGBColor(74, 83, 96)
    blue = RGBColor(31, 95, 168)
    panel = RGBColor(242, 244, 247)

    for idx, item in enumerate(slides):
        slide = prs.slides.add_slide(blank)
        title_box = slide.shapes.add_textbox(Inches(.48), Inches(.28), Inches(12.36), Inches(.62))
        tf = title_box.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.text = item.title
        p.font.name = "Aptos Display"
        if idx == 0:
            title_size = 31
        elif len(item.title) > 70:
            title_size = 18
        elif len(item.title) > 60:
            title_size = 21
        else:
            title_size = 28
        p.font.size = Pt(title_size)
        p.font.bold = True
        p.font.color.rgb = ink
        line = slide.shapes.add_shape(1, Inches(.48), Inches(.94), Inches(12.36), Inches(.025))
        line.fill.solid(); line.fill.fore_color.rgb = blue; line.line.fill.background()

        bullet_box = slide.shapes.add_textbox(Inches(.55), Inches(1.14), Inches(6.0), Inches(5.75))
        btf = bullet_box.text_frame
        btf.clear(); btf.word_wrap = True
        bullet_size = 15 if len(item.bullets) <= 5 else (10 if len(item.bullets) <= 14 else 8.5)
        bullet_space = 8 if len(item.bullets) <= 5 else (3 if len(item.bullets) <= 14 else 1)
        for j, bullet in enumerate(item.bullets):
            p = btf.paragraphs[0] if j == 0 else btf.add_paragraph()
            p.text = bullet
            p.level = 0
            p.font.name = "Aptos"
            p.font.size = Pt(bullet_size)
            p.font.color.rgb = ink
            p.space_after = Pt(bullet_space)

        mono_box = slide.shapes.add_textbox(Inches(6.75), Inches(1.15), Inches(6.05), Inches(5.72))
        mono_box.fill.solid(); mono_box.fill.fore_color.rgb = panel
        mono_box.line.color.rgb = RGBColor(216, 220, 227)
        mtf = mono_box.text_frame
        mtf.clear(); mtf.word_wrap = True; mtf.vertical_anchor = MSO_ANCHOR.TOP
        text = (("\n" if len(item.mono) > 6 else "\n\n").join(item.mono)) or "drperf"
        p = mtf.paragraphs[0]
        p.text = text
        p.font.name = "DejaVu Sans Mono"
        p.font.size = Pt(13)
        p.font.color.rgb = ink
        p.alignment = PP_ALIGN.LEFT

        num = slide.shapes.add_textbox(Inches(12.25), Inches(7.12), Inches(.65), Inches(.2))
        ntf = num.text_frame; ntf.clear(); np = ntf.paragraphs[0]
        np.text = f"{idx + 1} / {len(slides)}"; np.font.name = "DejaVu Sans Mono"; np.font.size = Pt(8); np.font.color.rgb = dim
        try:
            slide.notes_slide.notes_text_frame.text = item.notes
        except Exception:
            pass

    prs.save(PPTX_OUT)
    return True


def verify() -> tuple[list[str], list[str]]:
    cases_text = CASES.read_text()
    unmatched_formulas = [line for line in FORMULAS if line not in cases_text]
    unmatched_markers = []
    for path, line in MARKERS:
        p = ROOT / path if not path.startswith("/") else Path(path)
        if not p.exists() or line not in p.read_text():
            unmatched_markers.append(f"{path}: {line}")
    return unmatched_formulas, unmatched_markers


if __name__ == "__main__":
    missing_formula, missing_marker = verify()
    if missing_formula or missing_marker:
        raise SystemExit(f"verification failed before write: formulas={missing_formula!r}; markers={missing_marker!r}")
    build_html()
    pptx_ok = build_pptx()
    print(f"slides={len(slides)} pptx={pptx_ok}")
    for i, slide in enumerate(slides, 1):
        print(f"{i:02d} {slide.title}")
    print(f"unmatched_formula_lines={len(missing_formula)}")
    print(f"unmatched_marker_lines={len(missing_marker)}")
